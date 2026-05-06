from openpyxl import Workbook
from openpyxl.styles import Font

from src.schemas import CompetencyMap, FL1


def write_excel(comp_map: CompetencyMap, out_path: str) -> None:
    """Write CompetencyMap to .xlsx matching LU14 format."""
    wb = Workbook()

    _write_attribute_sheet(wb, comp_map)
    _write_per_fl1_prereq_sheets(wb, comp_map)
    _write_gesamt_sheet(wb, comp_map)

    wb.save(out_path)


def _write_attribute_sheet(wb: Workbook, comp_map: CompetencyMap) -> None:
    ws = wb.active
    ws.title = "Attribute"

    # Title row
    ws["B1"] = f"LU 01 - {comp_map.topic}"
    ws["B1"].font = Font(bold=True, size=14)

    # Header row 3
    headers = ["K_ID", "Competencies", "FL1_ID", "FL1", "FL2_ID", "FL2", "FL3_ID", "FL3"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = Font(bold=True)

    # Column widths
    widths = [8, 50, 8, 50, 8, 50, 10, 55]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    # Data rows starting at row 4
    row = 4
    for comp in comp_map.competencies:
        k_first_row = True
        for fl1 in comp.children:
            fl1_first_row = True
            for fl2 in fl1.children:
                fl2_first_row = True
                if fl2.children:
                    # Emit one row per FL3 leaf
                    for fl3 in fl2.children:
                        _write_row(ws, row, comp, fl1, fl2, k_first_row, fl1_first_row, fl2_first_row,
                                   fl3_id=fl3.id, fl3_desc=fl3.description)
                        k_first_row = False
                        fl1_first_row = False
                        fl2_first_row = False
                        row += 1
                    # Summary row for the FL2 (FL2 id+desc in G/H columns)
                    ws.cell(row=row, column=7, value=fl2.id)
                    ws.cell(row=row, column=8, value=fl2.description)
                    row += 1
                else:
                    # FL2 with no FL3 children: single row
                    _write_row(ws, row, comp, fl1, fl2, k_first_row, fl1_first_row, fl2_first_row,
                               fl3_id="", fl3_desc=fl2.description)
                    k_first_row = False
                    fl1_first_row = False
                    fl2_first_row = False
                    row += 1


def _write_row(ws, row, comp, fl1, fl2, k_first, fl1_first, fl2_first, fl3_id, fl3_desc):
    if k_first:
        ws.cell(row=row, column=1, value=comp.id)
        ws.cell(row=row, column=2, value=comp.description)
    if fl1_first:
        ws.cell(row=row, column=3, value=fl1.id)
        ws.cell(row=row, column=4, value=fl1.description)
    if fl2_first:
        ws.cell(row=row, column=5, value=fl2.id)
        ws.cell(row=row, column=6, value=fl2.description)
    ws.cell(row=row, column=7, value=fl3_id)
    ws.cell(row=row, column=8, value=fl3_desc)


def _write_per_fl1_prereq_sheets(wb: Workbook, comp_map: CompetencyMap) -> None:
    all_fl1s = _collect_fl1s(comp_map)
    prereq_set = {(e.prerequisite_id, e.dependent_id) for e in comp_map.prerequisites}

    for fl1 in all_fl1s:
        if not fl1.children:
            continue
        ws = wb.create_sheet(title=f"Prereq_{fl1.id}")
        fl2_list = fl1.children
        _write_matrix(ws, fl2_list, prereq_set)


def _write_gesamt_sheet(wb: Workbook, comp_map: CompetencyMap) -> None:
    ws = wb.create_sheet(title="Prerk_gesamt")
    all_fl2 = _collect_all_fl2(comp_map)
    prereq_set = {(e.prerequisite_id, e.dependent_id) for e in comp_map.prerequisites}
    _write_matrix(ws, all_fl2, prereq_set)


def _write_matrix(ws, fl2_list, prereq_set):
    """Write a prerequisite matrix. Row = prerequisite, Col = dependent."""
    # Header row: B1, C1, ... = FL2 descriptions
    for col_idx, fl2 in enumerate(fl2_list, start=2):
        ws.cell(row=1, column=col_idx, value=fl2.description)

    # Column A from row 2: FL2 descriptions
    for row_idx, fl2_row in enumerate(fl2_list, start=2):
        ws.cell(row=row_idx, column=1, value=fl2_row.description)

        for col_idx, fl2_col in enumerate(fl2_list, start=2):
            # Diagonal = 1
            if fl2_row.id == fl2_col.id:
                ws.cell(row=row_idx, column=col_idx, value=1)
            elif (fl2_row.id, fl2_col.id) in prereq_set:
                ws.cell(row=row_idx, column=col_idx, value=1)


def _collect_fl1s(comp_map: CompetencyMap) -> list[FL1]:
    result = []
    for comp in comp_map.competencies:
        result.extend(comp.children)
    return result


def _collect_all_fl2(comp_map: CompetencyMap):
    from src.schemas import FL2
    result: list[FL2] = []
    for comp in comp_map.competencies:
        for fl1 in comp.children:
            result.extend(fl1.children)
    return result
