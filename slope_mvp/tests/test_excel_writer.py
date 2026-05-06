import os
import tempfile

from openpyxl import load_workbook

from src.schemas import Competency, FL1, FL2, FL3, CompetencyMap, PrereqEdge
from src.excel_writer import write_excel


def _make_fixture() -> CompetencyMap:
    return CompetencyMap(
        topic="Slope of Lines",
        competencies=[
            Competency(
                id="K1",
                description="understand slope as a ratio of two quantities",
                children=[
                    FL1(
                        id="F1",
                        description="determine the equation of a line from a graph",
                        children=[
                            FL2(
                                id="F1.0",
                                description="construct a readable slope triangle",
                                children=[
                                    FL3(id="F1.0.1", description="slope triangle when m>0", parametric_condition="m>0"),
                                    FL3(id="F1.0.2", description="slope triangle when m<0", parametric_condition="m<0"),
                                ],
                            ),
                            FL2(
                                id="F1.1",
                                description="read the slope using a slope triangle",
                                children=[],
                            ),
                            FL2(
                                id="F1.2",
                                description="determine the equation y = mx",
                                children=[],
                            ),
                        ],
                    ),
                ],
            ),
            Competency(
                id="K2",
                description="interpret slope as rate of change",
                children=[
                    FL1(
                        id="F2",
                        description="plot a line given its equation",
                        children=[
                            FL2(
                                id="F2.0",
                                description="read the y-intercept",
                                children=[
                                    FL3(id="F2.0.1", description="y-intercept when b>0", parametric_condition="b>0"),
                                    FL3(id="F2.0.2", description="y-intercept when b<0", parametric_condition="b<0"),
                                ],
                            ),
                        ],
                    ),
                ],
            ),
        ],
        prerequisites=[
            PrereqEdge(prerequisite_id="F1.0", dependent_id="F1.1", rationale="need triangle to read slope"),
            PrereqEdge(prerequisite_id="F1.1", dependent_id="F1.2", rationale="need slope reading to form equation"),
        ],
    )


def test_write_excel_creates_file():
    cm = _make_fixture()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        path = f.name

    try:
        write_excel(cm, path)
        assert os.path.exists(path)
        wb = load_workbook(path)

        # Sheet names
        assert "Attribute" in wb.sheetnames
        assert "Prereq_F1" in wb.sheetnames
        assert "Prereq_F2" in wb.sheetnames
        assert "Prerk_gesamt" in wb.sheetnames

        # Attribute sheet structure
        ws = wb["Attribute"]
        assert ws["B1"].value == "LU 01 - Slope of Lines"
        assert ws.cell(row=3, column=1).value == "K_ID"
        assert ws.cell(row=4, column=1).value == "K1"
        assert ws.cell(row=4, column=2).value == "understand slope as a ratio of two quantities"

        # Prerk_gesamt: should be a 4x4 matrix (4 FL2 nodes total)
        ws_g = wb["Prerk_gesamt"]
        # Row 1 headers start at B1
        assert ws_g.cell(row=1, column=2).value is not None
        # Diagonal should be 1
        for i in range(2, 6):
            assert ws_g.cell(row=i, column=i).value == 1

        # Prereq_F1: 3x3 matrix (F1.0, F1.1, F1.2)
        ws_f1 = wb["Prereq_F1"]
        # F1.0 → F1.1 should be marked
        # F1.0 is row 2, F1.1 is col 3
        assert ws_f1.cell(row=2, column=3).value == 1
    finally:
        os.unlink(path)
